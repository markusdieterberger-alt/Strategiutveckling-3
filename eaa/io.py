"""Strict CSV/XLSX import. Unsupported layouts fail instead of guessing P&L."""

from __future__ import annotations

import csv
from dataclasses import replace
from datetime import date, datetime
import math
from pathlib import Path
from zoneinfo import ZoneInfo

from .funded import Sample, Trade


def number(value):
    raw = str(value).strip().replace("\u2212", "-")
    # English export format only. Locale conversion requires an explicit map.
    raw = raw.replace("$", "").replace(" USD", "")
    if "," in raw:
        import re
        if not re.fullmatch(r"-?\d{1,3}(,\d{3})+(\.\d+)?", raw):
            raise ValueError(f"Ambiguous numeric locale: {value!r}")
        raw = raw.replace(",", "")
    result = float(raw)
    if not math.isfinite(result):
        raise ValueError("Non-finite numeric value")
    return result


def timestamp(value, timezone=None):
    at = value if isinstance(value, datetime) else datetime.fromisoformat(str(value).strip().replace("Z", "+00:00"))
    if at.tzinfo is None:
        if timezone is None:
            raise ValueError("Naive export time requires --export-timezone")
        zone = ZoneInfo(timezone)
        a, b = at.replace(tzinfo=zone, fold=0), at.replace(tzinfo=zone, fold=1)
        if a.utcoffset() != b.utcoffset():
            raise ValueError("Ambiguous/nonexistent DST time: provide an explicit UTC offset")
        at = a
    return at


def read_rows(path, sheet=None):
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        try:
            import openpyxl
        except ImportError as error:
            raise ValueError("XLSX needs optional openpyxl; export the Trades sheet as CSV") from error
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        if sheet is None:
            matches = [s for s in book.sheetnames if s.lower() in {"trades", "list of trades"}]
            if len(matches) != 1:
                raise ValueError(f"Select --sheet explicitly from {book.sheetnames}")
            sheet = matches[0]
        values = iter(book[sheet].values)
        header = [str(x).strip() for x in next(values)]
        rows = [dict(zip(header, row)) for row in values if any(x is not None for x in row)]
        book.close()
    else:
        with path.open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            header, rows = reader.fieldnames or [], list(reader)
    if not rows or len(header) != len(set(header)):
        raise ValueError("Empty input or duplicate column names")
    if any(None in r for r in rows):
        raise ValueError("Malformed CSV row")
    return rows


def load_trades(path, timezone=None, mapping=None, sheet=None):
    rows = read_rows(path, sheet)
    if mapping:
        rows = [{mapping.get(k, k): v for k, v in row.items()} for row in rows]
    header = rows[0]
    if {"trade_id", "entry_time", "exit_time", "side", "entry_price", "exit_price"} <= header.keys():
        trades = []
        for r in rows:
            side = {"long": 1, "short": -1}.get(str(r["side"]).lower())
            qty = number(r.get("contracts") or 1)
            if not qty.is_integer():
                raise ValueError("Fractional contracts unsupported")
            trades.append(Trade(str(r["trade_id"]), timestamp(r["entry_time"], timezone), timestamp(r["exit_time"], timezone), side, number(r["entry_price"]), number(r["exit_price"]), int(qty), number(r["stop_points"]) if r.get("stop_points") else None, number(r["mae_points"]) if r.get("mae_points") else None))
    else:
        def column(options):
            found = [k for k in options if k in header]
            if len(found) != 1:
                raise ValueError(f"Need exactly one column from {options}; use --column-map")
            return found[0]
        idcol = column(["Trade #", "Trade number", "trade_id"])
        typecol = column(["Type", "type"])
        timecol = column(["Date/Time", "Date and time", "time"])
        pricecol = column(["Price USD", "Price", "price"])
        qtycol = column(["Position size (qty)", "Contracts", "Quantity", "contracts"])
        pairs = {}
        for row in rows:
            ident = str(row[idcol])
            kind = str(row[typecol]).strip().lower()
            if kind not in {"entry long", "entry short", "exit long", "exit short"}:
                raise ValueError(f"Unsupported row type {kind!r}; open/partial positions need reconciliation")
            role, direction = kind.split()
            pair = pairs.setdefault(ident, {})
            if role in pair:
                raise ValueError("Multiple entries/exits per id unsupported")
            pair[role] = (row, direction)
        trades = []
        for ident, pair in pairs.items():
            if set(pair) != {"entry", "exit"}:
                raise ValueError(f"Unpaired/open trade {ident}")
            (entry, side), (exit_row, exit_side) = pair["entry"], pair["exit"]
            qty, exit_qty = number(entry[qtycol]), number(exit_row[qtycol])
            if side != exit_side or qty != exit_qty or not qty.is_integer():
                raise ValueError("Direction/quantity mismatch or partial exit")
            trades.append(Trade(ident, timestamp(entry[timecol], timezone), timestamp(exit_row[timecol], timezone), 1 if side == "long" else -1, number(entry[pricecol]), number(exit_row[pricecol]), int(qty)))
    return sorted(trades, key=lambda t: (t.entry, t.exit, t.id))


def attach_paths(trades, path, timezone=None):
    known = {t.id for t in trades}
    grouped = {ident: [] for ident in known}
    for row in read_rows(path):
        ident = str(row["trade_id"])
        if ident not in known:
            raise ValueError(f"Unknown path trade {ident}")
        grouped[ident].append(Sample(timestamp(row["time"], timezone), number(row["unrealized_points"])))
    if any(not grouped[t.id] for t in trades):
        raise ValueError("Intraday path missing for at least one trade")
    return [replace(t, samples=tuple(grouped[t.id]), quality="SAMPLED_INTRADAY_PATH") for t in trades]


def load_calendar(path):
    days = [date.fromisoformat(str(r["session_date"])) for r in read_rows(path)]
    if days != sorted(set(days)):
        raise ValueError("Calendar must be ordered with unique session dates")
    return days
